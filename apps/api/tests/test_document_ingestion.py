"""Tests for document ingestion service."""

import pytest
from pathlib import Path

from app.ingestion.service import DocumentIngestionService, DocumentIngestionConfig
from app.ingestion.extractors import (
    ContentType,
    TextExtractor,
    XLSXExtractor,
    DocxExtractor,
)
from app.ingestion.markdown_builder import build_simple_markdown


class TestTextExtractor:
    """Tests for text extractor."""

    def test_extract_file(self, tmp_path):
        """Test extracting from a text file."""
        test_file = tmp_path / "test.txt"
        test_file.write_text("Hello, world!\nThis is a test.")

        extractor = TextExtractor()
        result = extractor.extract(test_file)

        assert result.content_type == ContentType.TEXT
        assert "Hello, world!" in result.text
        assert result.source_filename == "test.txt"

    def test_extract_bytes(self):
        """Test extracting from bytes."""
        content = b"Plain text content\nWith multiple lines"
        extractor = TextExtractor()
        result = extractor.extract_bytes(content, "test.txt")

        assert result.content_type == ContentType.TEXT
        assert "Plain text content" in result.text


class TestXLSXExtractor:
    """Tests for XLSX extractor."""

    def test_extract_simple_spreadsheet(self, tmp_path):
        """Test extracting a simple spreadsheet."""
        import openpyxl

        test_file = tmp_path / "test.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Name", "Value", "Status"])
        ws.append(["Item 1", "100", "Active"])
        ws.append(["Item 2", "200", "Inactive"])
        wb.save(test_file)

        extractor = XLSXExtractor()
        result = extractor.extract(test_file)

        assert result.content_type == ContentType.XLSX
        assert len(result.tables) == 1
        table = result.tables[0]
        assert table.headers == ["Name", "Value", "Status"]
        assert len(table.rows) == 2
        assert table.rows[0] == ["Item 1", "100", "Active"]


class TestDocxExtractor:
    """Tests for DOCX extractor."""

    def test_extract_simple_document(self, tmp_path):
        """Test extracting a simple Word document."""
        from docx import Document

        test_file = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("This is a test document.")
        doc.add_paragraph("It has multiple paragraphs.")
        doc.save(test_file)

        extractor = DocxExtractor()
        result = extractor.extract(test_file)

        assert result.content_type == ContentType.DOCX
        assert "This is a test document" in result.text
        assert "multiple paragraphs" in result.text

    def test_extract_document_with_table(self, tmp_path):
        """Test extracting a Word document with a table."""
        from docx import Document

        test_file = tmp_path / "test_table.docx"
        doc = Document()
        doc.add_paragraph("Table below:")
        table = doc.add_table(rows=3, cols=2)
        table.rows[0].cells[0].text = "Header 1"
        table.rows[0].cells[1].text = "Header 2"
        table.rows[1].cells[0].text = "Row 1 Col 1"
        table.rows[1].cells[1].text = "Row 1 Col 2"
        table.rows[2].cells[0].text = "Row 2 Col 1"
        table.rows[2].cells[1].text = "Row 2 Col 2"
        doc.save(test_file)

        extractor = DocxExtractor()
        result = extractor.extract(test_file)

        assert len(result.tables) == 1
        table = result.tables[0]
        assert table.headers == ["Header 1", "Header 2"]
        assert len(table.rows) == 2


class TestMarkdownBuilder:
    """Tests for markdown builder."""

    def test_build_simple_markdown(self):
        """Test building simple markdown from extracted content."""
        from app.ingestion.extractors import ExtractedContent, Table

        extracted = ExtractedContent(
            text="This is the main text content.",
            tables=[
                Table(headers=["Col1", "Col2"], rows=[["A", "B"], ["C", "D"]])
            ],
            content_type=ContentType.TEXT,
            source_filename="test.txt",
        )

        markdown = build_simple_markdown(extracted)

        assert "This is the main text content" in markdown
        assert "Col1" in markdown
        assert "Col2" in markdown
        assert "A" in markdown
        assert "B" in markdown


class TestDocumentIngestionService:
    """Tests for the main document ingestion service."""

    @pytest.fixture
    def service(self):
        """Create a service instance for testing."""
        config = DocumentIngestionConfig(
            cache_dir=Path(".cache/test_ingestion"),
            enable_vision_fallback=False,  # Disable for tests
        )
        service = DocumentIngestionService(config)
        yield service
        service.shutdown()

    def test_ingest_text(self, service):
        """Test ingesting plain text."""
        result = service.ingest_text("This is test content.\nWith multiple lines.")

        assert "This is test content" in result.content
        assert result.metadata["extractor"] == "TextExtractor"

    def test_ingest_generic_text_source(self, service):
        """Test the simple generic ingestion entry point used by callers."""
        result = service.ingest("This is pasted text from a user copy\nwith a second line.")

        assert "This is pasted text" in result.content
        assert result.metadata["extractor"] == "TextExtractor"

    def test_ingest_txt_file(self, service, tmp_path):
        """Test ingesting a text file."""
        test_file = tmp_path / "sample.txt"
        test_file.write_text("File content here.")

        result = service.ingest_file(test_file)

        assert "File content here" in result.content
        assert result.source_filename == "sample.txt" or "sample.txt" in str(result.metadata)

    def test_ingest_xlsx_file(self, service, tmp_path):
        """Test ingesting an Excel file."""
        import openpyxl

        test_file = tmp_path / "data.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Product", "Price", "Qty"])
        ws.append(["Widget A", "10.00", "5"])
        ws.append(["Widget B", "20.00", "3"])
        wb.save(test_file)

        result = service.ingest_file(test_file)

        assert "Product" in result.content
        assert "Widget A" in result.content
        assert "10.00" in result.content

    def test_ingest_docx_file(self, service, tmp_path):
        """Test ingesting a Word document."""
        from docx import Document

        test_file = tmp_path / "doc.docx"
        doc = Document()
        doc.add_paragraph("Shipping Document")
        doc.add_paragraph("Bill of Lading: BL-12345")
        doc.save(test_file)

        result = service.ingest_file(test_file)

        assert "Shipping Document" in result.content
        assert "BL-12345" in result.content

    def test_cache_works(self, service, tmp_path):
        """Test that caching works."""
        test_file = tmp_path / "cache_test.txt"
        test_file.write_text("Cache test content.")

        # First ingestion
        result1 = service.ingest_file(test_file)

        # Second ingestion (should hit cache)
        result2 = service.ingest_file(test_file)

        assert result1.content == result2.content
        stats = service.get_cache_stats()
        assert stats["entry_count"] > 0


# Integration tests with real sample files (if available)
class TestIntegrationWithSampleFiles:
    """Integration tests with actual sample files from sdoc-data."""

    @pytest.fixture
    def service(self):
        config = DocumentIngestionConfig(
            cache_dir=Path(".cache/test_ingestion"),
            enable_vision_fallback=False,
        )
        service = DocumentIngestionService(config)
        yield service
        service.shutdown()

    def test_sample_txt_attachment(self, service):
        """Test with a real TXT attachment from sample data."""
        sample_file = Path("sdoc-data/attachments/email_001_BL.txt")
        if not sample_file.exists():
            pytest.skip("Sample file not found")

        result = service.ingest_file(sample_file)
        assert len(result.content) > 0
        assert "email_001_BL.txt" in str(result.metadata) or result.metadata.get("source_filename") == "email_001_BL.txt"

    def test_sample_xlsx_attachment(self, service):
        """Test with a real XLSX attachment from sample data."""
        sample_file = Path("sdoc-data/attachments/email_005_BL.xlsx")
        if not sample_file.exists():
            pytest.skip("Sample file not found")

        result = service.ingest_file(sample_file)
        assert len(result.content) > 0
        # Should have extracted table data
        assert "|" in result.content  # Markdown table syntax

    def test_ingest_all_attachments(self, service):
        """Ingest every attachment in sdoc-data and verify non‑empty output."""
        attachments_dir = Path("sdoc-data/attachments")
        if not attachments_dir.is_dir():
            pytest.skip("Attachments directory not found")
        for file_path in sorted(attachments_dir.iterdir()):
            if not file_path.is_file():
                continue
            result = service.ingest_file(file_path)
            assert result.content, f"Empty content for {file_path.name}"


if __name__ == "__main__":
    pytest.main([__file__, "-v"])