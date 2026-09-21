"""DOCX (Word) document extractor."""

import io
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table as DocxTable

from .base import ContentType, DocumentExtractor, ExtractedContent, Image, Table


class DocxExtractor(DocumentExtractor):
    """Extractor for Word documents (.docx)."""

    supported_types = [ContentType.DOCX]

    def extract(self, file_path: Path) -> ExtractedContent:
        """Extract content from a DOCX file."""
        doc = Document(str(file_path))
        return self._extract_document(doc, file_path)

    def extract_bytes(self, content: bytes, filename: str) -> ExtractedContent:
        """Extract content from raw bytes."""
        doc = Document(io.BytesIO(content))
        return self._extract_document(doc, Path(filename))

    def _extract_document(self, doc: DocxDocument, source: Path) -> ExtractedContent:
        """Extract all content from a Word document."""
        text_parts = []
        tables = []
        images = []
        metadata = {"paragraph_count": 0, "table_count": 0, "image_count": 0}

        # Extract paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
                metadata["paragraph_count"] += 1

        # Extract tables
        for i, table in enumerate(doc.tables):
            extracted_table = self._extract_table(table, i)
            if extracted_table:
                tables.append(extracted_table)
                text_parts.append(f"## Table {i + 1}")
                text_parts.append(extracted_table.to_markdown())
                metadata["table_count"] += 1

        # Extract images (from relationships)
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                try:
                    image_data = rel.target_part.blob
                    image = Image(
                        data=image_data,
                        mime_type=rel.target_part.content_type,
                        alt_text=f"Image from {source.name}",
                    )
                    images.append(image)
                    metadata["image_count"] += 1
                except Exception:
                    # Skip images that can't be extracted
                    pass

        return ExtractedContent(
            text="\n\n".join(text_parts),
            tables=tables,
            images=images,
            content_type=ContentType.DOCX,
            source_filename=source.name,
            metadata=metadata,
        )

    def _extract_table(self, table: DocxTable, table_index: int) -> Table | None:
        """Extract a Word table as a Table object."""
        rows = []
        headers = None

        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]

            # Skip completely empty rows
            if not any(cells):
                continue

            # The first *non-empty* row is the header (see xlsx extractor).
            if headers is None:
                headers = cells
            else:
                rows.append(cells)

        if not headers and not rows:
            return None

        return Table(
            headers=headers or [],
            rows=rows,
        )
