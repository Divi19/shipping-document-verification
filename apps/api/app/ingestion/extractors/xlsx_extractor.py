"""Excel (XLSX) extractor."""

import hashlib
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from .base import ContentType, DocumentExtractor, ExtractedContent, Table


class XLSXExtractor(DocumentExtractor):
    """Extractor for Excel spreadsheets."""

    supported_types = [ContentType.XLSX]

    def extract(self, file_path: Path) -> ExtractedContent:
        """Extract content from an XLSX file."""
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        return self._extract_workbook(workbook, file_path)

    def extract_bytes(self, content: bytes, filename: str) -> ExtractedContent:
        """Extract content from raw bytes."""
        import io
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        return self._extract_workbook(workbook, Path(filename))

    def _extract_workbook(self, workbook: openpyxl.Workbook, source: Path) -> ExtractedContent:
        """Extract all sheets from workbook."""
        tables = []
        metadata = {
            "sheet_names": workbook.sheetnames,
            "sheet_count": len(workbook.sheetnames),
        }

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            table = self._extract_sheet(sheet, sheet_name)
            if table:
                tables.append(table)

        # Build text representation from tables
        text_parts = []
        for table in tables:
            if table.sheet_name:
                text_parts.append(f"## Sheet: {table.sheet_name}")
            text_parts.append(table.to_markdown())

        workbook.close()

        return ExtractedContent(
            text="\n\n".join(text_parts),
            tables=tables,
            content_type=ContentType.XLSX,
            source_filename=source.name,
            metadata=metadata,
        )

    def _extract_sheet(self, sheet: Worksheet, sheet_name: str) -> Optional[Table]:
        """Extract a single sheet as a table."""
        rows = []
        headers = None

        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            # Skip completely empty rows
            if all(cell is None for cell in row):
                continue

            # Convert all cells to strings
            str_row = [str(cell) if cell is not None else "" for cell in row]

            if i == 0:
                headers = str_row
            else:
                rows.append(str_row)

        if not headers and not rows:
            return None

        return Table(
            headers=headers or [],
            rows=rows,
            sheet_name=sheet_name,
        )